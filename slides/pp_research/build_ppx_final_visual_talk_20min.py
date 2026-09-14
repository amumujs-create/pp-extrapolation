#!/usr/bin/env python3
"""Build the image-led PP-X Final 20-minute presentation."""
from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from matplotlib import font_manager
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "ppt" / "PP-X_Final_Visual_Talk_20min.pptx"
FIG = ROOT / "ppt" / "_final_visual_assets"

W, H = Inches(13.333), Inches(7.5)
INK = RGBColor(26, 35, 47)
MUTED = RGBColor(91, 103, 119)
BLUE = RGBColor(31, 107, 181)
CYAN = RGBColor(45, 157, 184)
ORANGE = RGBColor(230, 126, 52)
RED = RGBColor(190, 65, 65)
GREEN = RGBColor(45, 145, 104)
WHITE = RGBColor(255, 255, 255)
PALE_BLUE = RGBColor(233, 243, 251)
PALE_ORANGE = RGBColor(253, 241, 230)
PALE_GREEN = RGBColor(231, 246, 239)
PALE_GRAY = RGBColor(244, 246, 248)


def rgb(c: RGBColor) -> str:
    return f"#{c[0]:02x}{c[1]:02x}{c[2]:02x}"


def choose_font() -> str:
    candidates = ["Apple SD Gothic Neo", "Noto Sans CJK KR", "NanumGothic", "Arial Unicode MS"]
    available = {f.name for f in font_manager.fontManager.ttflist}
    return next((name for name in candidates if name in available), "DejaVu Sans")


FONT = choose_font()
plt.rcParams.update({"font.family": FONT, "axes.unicode_minus": False})


def blank(prs: Presentation, dark: bool = False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = INK if dark else WHITE
    return s


def text(s, x, y, w, h, value, size=18, color=INK, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    r = p.add_run()
    r.text = value
    r.font.name = "Apple SD Gothic Neo"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def box(s, x, y, w, h, fill=PALE_GRAY, line=None, radius=True):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line or fill
    sh.line.width = Pt(1.2)
    return sh


def line(s, x1, y1, x2, y2, color=MUTED, width=2.0, arrow=False):
    sh = s.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    sh.line.color.rgb = color
    sh.line.width = Pt(width)
    if arrow:
        # python-pptx does not expose arrowheads consistently across releases.
        # A small terminal marker preserves the direction without private XML.
        dot(s, x2 - .05, y2 - .05, .1, color)
    return sh


def dot(s, x, y, d, fill):
    sh = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def title(s, value, sub=None, n=None):
    text(s, .55, .25, 11.8, .45, value, 25, INK, True)
    if sub:
        text(s, .58, .72, 11.7, .28, sub, 10.5, MUTED)
    if n:
        text(s, 12.25, .32, .5, .3, f"{n:02}", 10, MUTED, True, PP_ALIGN.RIGHT)
    line(s, .55, 1.08, 12.78, 1.08, RGBColor(223, 227, 232), 1)


def pill(s, x, y, w, value, fill, color=WHITE, size=11):
    box(s, x, y, w, .36, fill, fill)
    text(s, x, y, w, .36, value, size, color, True, PP_ALIGN.CENTER)


def formula(s, x, y, w, h, value, caption, fill=PALE_BLUE, accent=BLUE):
    box(s, x, y, w, h, fill, accent)
    text(s, x + .2, y + .12, w - .4, h * .55, value, 20, accent, True, PP_ALIGN.CENTER)
    text(s, x + .2, y + h * .62, w - .4, h * .26, caption, 10.5, INK, False, PP_ALIGN.CENTER)


def save_chart(path: Path, kind: str):
    FIG.mkdir(parents=True, exist_ok=True)
    if kind == "scores":
        names = ["HUST", "Sunwoda", "N-CMAPSS", "Virkler", "RWTH", "MATR-b2", "MICH", "NASA", "MATR19"]
        vals = np.array([.958, .939, .937, .888, .878, .862, .751, .584, .466])
        fig, ax = plt.subplots(figsize=(10.7, 4.6))
        colors = [rgb(BLUE) if v >= .85 else rgb(CYAN) for v in vals]
        bars = ax.barh(np.arange(9), vals, color=colors, height=.66)
        ax.set_yticks(np.arange(9), names)
        ax.invert_yaxis()
        ax.set_xlim(0, 1.02)
        ax.set_xlabel("pooled R²")
        ax.grid(axis="x", color="#dfe3e8", linewidth=.8)
        ax.spines[["top", "right", "left"]].set_visible(False)
        ax.bar_label(bars, fmt="%.3f", padding=5, fontsize=10, fontweight="bold")
        fig.tight_layout()
    elif kind == "equal":
        names = ["HUST", "Sunwoda", "N-CMAPSS", "Virkler", "RWTH", "MATR-b2", "MICH", "NASA", "MATR19"]
        ppx = np.array([.958, .939, .937, .888, .878, .862, .751, .584, .466])
        base = np.array([.955, .838, .932, .890, .732, .813, -.686, .583, .342])
        fig, ax = plt.subplots(figsize=(10.8, 4.6))
        yy = np.arange(9)
        ax.hlines(yy, base, ppx, color="#cbd2d9", linewidth=3)
        ax.scatter(base, yy, s=56, color=rgb(MUTED), label="30-candidate strongest baseline", zorder=3)
        ax.scatter(ppx, yy, s=72, color=rgb(BLUE), label="PP-X Final", zorder=4)
        ax.set_yticks(yy, names)
        ax.invert_yaxis()
        ax.axvline(0, color="#dfe3e8", linewidth=1)
        ax.set_xlim(-.8, 1.05)
        ax.set_xlabel("pooled R²")
        ax.grid(axis="x", color="#edf0f2", linewidth=.8)
        ax.spines[["top", "right", "left"]].set_visible(False)
        ax.legend(loc="lower right", frameon=False, ncol=2, fontsize=9)
        fig.tight_layout()
    elif kind == "core":
        names = ["Sunwoda", "RWTH", "MICH"]
        affine = np.array([.281, .659, -3.343])
        final = np.array([.939, .878, .468])
        x = np.arange(3)
        fig, ax = plt.subplots(figsize=(10.5, 4.5))
        b1 = ax.bar(x - .18, affine, .36, color=rgb(MUTED), label="affine quotient only")
        b2 = ax.bar(x + .18, final, .36, color=rgb(BLUE), label="BQ + nonlinear residual")
        ax.axhline(0, color="#aeb6bf", linewidth=1)
        ax.set_xticks(x, names)
        ax.set_ylabel("pooled R²")
        ax.set_ylim(-3.7, 1.15)
        ax.grid(axis="y", color="#e5e8eb", linewidth=.8)
        ax.spines[["top", "right", "left"]].set_visible(False)
        ax.bar_label(b1, fmt="%.3f", padding=3, fontsize=9)
        ax.bar_label(b2, fmt="%.3f", padding=3, fontsize=9, fontweight="bold")
        ax.legend(frameon=False, loc="lower left")
        fig.tight_layout()
    elif kind == "executor":
        labels = ["HUST transport", "MATR-b2 transport", "MICH dual-scale",
                  "Sunwoda bound", "RWTH bound", "RWTH dual-scale", "MICH fixed bound"]
        vals = np.array([.128, .187, .283, .221, .090, -.037, -.291])
        colors = [rgb(GREEN) if v > 0 else rgb(RED) for v in vals]
        fig, ax = plt.subplots(figsize=(10.6, 4.5))
        bars = ax.barh(np.arange(len(vals)), vals, color=colors, height=.62)
        ax.set_yticks(np.arange(len(vals)), labels)
        ax.invert_yaxis()
        ax.axvline(0, color="#8f99a5", linewidth=1.2)
        ax.set_xlim(-.34, .34)
        ax.set_xlabel("matched ΔR² (on − off)")
        ax.grid(axis="x", color="#e5e8eb", linewidth=.8)
        ax.spines[["top", "right", "left"]].set_visible(False)
        ax.bar_label(bars, fmt="%+.3f", padding=4, fontsize=10, fontweight="bold")
        fig.tight_layout()
    else:
        raise ValueError(kind)
    fig.savefig(path, dpi=190, transparent=True, bbox_inches="tight")
    plt.close(fig)


def add_picture(s, path, x=.7, y=1.3, w=11.9, h=5.55):
    s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


def build():
    FIG.mkdir(parents=True, exist_ok=True)
    charts = {}
    for kind in ("scores", "equal", "core", "executor"):
        charts[kind] = FIG / f"{kind}.png"
        save_chart(charts[kind], kind)

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    slide_no = 0

    # 1 Cover
    s = blank(prs, True)
    pill(s, .72, .62, 2.25, "20 MIN · PP-X FINAL", BLUE)
    text(s, .75, 1.42, 11.8, 1.0, "외삽을 안전하게 만드는\nprior–residual 실행 규칙", 34, WHITE, True)
    text(s, .78, 3.8, 10.4, .6, "PP-X Final 게이트 · 수식 · 성능 · ablation · 통계검정", 17, RGBColor(207, 219, 231))
    line(s, .78, 4.65, 4.0, 4.65, BLUE, 5)
    text(s, .78, 5.3, 6.0, .5, "박사과정 박진서", 16, WHITE, True)
    text(s, .78, 5.82, 7.5, .35, "Validation-Approved Prior-Residual Extrapolation", 11, RGBColor(180, 195, 210))
    text(s, 11.4, 6.75, 1.2, .3, "01 / 16", 10, RGBColor(180, 195, 210), True, PP_ALIGN.RIGHT)

    # 2 Why extrapolation
    slide_no = 2
    s = blank(prs)
    title(s, "왜 외삽인가", "고장에 가까운 구간은 학습 support 밖에 놓이기 쉽다.", slide_no)
    line(s, 1.0, 4.95, 12.2, 4.95, MUTED, 2)
    line(s, 4.8, 1.6, 4.8, 5.55, MUTED, 1.5)
    text(s, 1.0, 5.18, 3.4, .35, "학습 구간", 13, MUTED, True, PP_ALIGN.CENTER)
    text(s, 5.1, 5.18, 6.6, .35, "외삽 구간 · EOL 방향", 13, ORANGE, True, PP_ALIGN.CENTER)
    for i, yy in enumerate([2.0, 2.45, 2.9]):
        line(s, 1.2, yy, 4.55, yy + .9, BLUE, 3)
    line(s, 4.8, 3.05, 11.6, 4.7, GREEN, 4)
    line(s, 4.8, 3.05, 11.6, 2.0, RED, 3)
    text(s, 8.2, 4.45, 3.2, .4, "물리적으로 일관된 tail", 12, GREEN, True)
    text(s, 8.2, 1.7, 3.4, .5, "보간은 맞아도 tail은 되살아날 수 있음", 12, RED, True)
    box(s, .9, 6.02, 11.6, .72, PALE_ORANGE, ORANGE)
    text(s, 1.15, 6.12, 11.1, .5, "핵심: 관측 데이터만으로 support 밖 함수값은 식별되지 않는다 → 외삽 방향을 prior와 contract로 제한", 15, INK, True, PP_ALIGN.CENTER)

    # 3 PP-X to PAE
    slide_no += 1
    s = blank(prs)
    title(s, "연구 방향 — PP-X에서 PAE로", "검증된 식의 강도에 따라 두 층을 분리한다.", slide_no)
    box(s, .8, 1.55, 5.25, 4.35, PALE_BLUE, BLUE)
    pill(s, 1.12, 1.85, 1.45, "현재", BLUE)
    text(s, 1.12, 2.35, 4.5, .55, "PP-X", 28, BLUE, True)
    text(s, 1.12, 3.05, 4.45, 1.35, "약한 prior\n+ bounded residual\n+ validation 승인", 18, INK, True)
    text(s, 1.12, 4.8, 4.4, .62, "식이 약하거나 경계·방향만 아는 경우", 12, MUTED)
    line(s, 6.1, 3.7, 7.15, 3.7, ORANGE, 4, True)
    text(s, 6.15, 3.12, .95, .4, "확장", 11, ORANGE, True, PP_ALIGN.CENTER)
    box(s, 7.25, 1.55, 5.25, 4.35, PALE_ORANGE, ORANGE)
    pill(s, 7.57, 1.85, 1.6, "후속", ORANGE)
    text(s, 7.57, 2.35, 4.5, .55, "PAE", 28, ORANGE, True)
    text(s, 7.57, 3.05, 4.45, 1.35, "출처가 고정된 식\n+ 적용 가능성 검증\n+ residual 보정", 18, INK, True)
    text(s, 7.57, 4.8, 4.4, .62, "정당화된 방정식이 있는 경우에만", 12, MUTED)
    text(s, 1.0, 6.28, 11.3, .4, "식이 검증되지 않으면 PAE를 강제하지 않고 PP-X로 되돌린다.", 15, INK, True, PP_ALIGN.CENTER)

    # 4 Contribution
    slide_no += 1
    s = blank(prs)
    title(s, "PP-X가 고정하는 세 가지", "Declare → Learn → Approve", slide_no)
    cards = [
        (.75, "C1", "Contract", "경계 · causal X ·\n허용 executor · fallback", BLUE, PALE_BLUE),
        (4.55, "C2", "Prior–residual", "frozen tail prior +\n제한된 nonlinear correction", CYAN, RGBColor(230, 247, 249)),
        (8.35, "C3", "Validation approval", "gain · unit wins ·\nworst-unit risk", ORANGE, PALE_ORANGE),
    ]
    for x, code, name, body, accent, fill in cards:
        box(s, x, 1.55, 3.45, 4.65, fill, accent)
        dot(s, x + .28, 1.87, .52, accent)
        text(s, x + .28, 1.87, .52, .52, code, 11, WHITE, True, PP_ALIGN.CENTER)
        text(s, x + .3, 2.62, 2.9, .5, name, 19, accent, True)
        text(s, x + .3, 3.35, 2.9, 1.4, body, 15, INK, True)
        text(s, x + .3, 5.18, 2.9, .5, ["무엇을 믿나?", "얼마나 수정하나?", "써도 되나?"][cards.index((x, code, name, body, accent, fill))], 11, MUTED)

    # 5 Final route
    slide_no += 1
    s = blank(prs)
    title(s, "PP-X Final prior route", "Final은 known_boundary 하나만 본다. OOF/group/mode 사다리는 실행하지 않는다.", slide_no)
    box(s, 4.65, 1.35, 4.0, .75, INK, INK)
    text(s, 4.65, 1.35, 4.0, .75, "known_boundary ?", 20, WHITE, True, PP_ALIGN.CENTER)
    line(s, 6.65, 2.1, 6.65, 2.6, INK, 3)
    line(s, 2.8, 2.6, 10.5, 2.6, INK, 3)
    line(s, 2.8, 2.6, 2.8, 3.05, BLUE, 3, True)
    line(s, 10.5, 2.6, 10.5, 3.05, ORANGE, 3, True)
    pill(s, 1.9, 2.73, 1.8, "TRUE", BLUE)
    pill(s, 9.6, 2.73, 1.8, "FALSE", ORANGE)
    box(s, .85, 3.25, 4.9, 2.3, PALE_BLUE, BLUE)
    text(s, 1.15, 3.52, 4.3, .45, "BQ prior", 22, BLUE, True, PP_ALIGN.CENTER)
    text(s, 1.15, 4.12, 4.3, .8, "margin × positive quotient\n경계에서 RUL = 0", 15, INK, True, PP_ALIGN.CENTER)
    pill(s, 2.35, 5.02, 1.9, "PRIOR ON", BLUE)
    box(s, 7.55, 3.25, 4.9, 2.3, PALE_ORANGE, ORANGE)
    text(s, 7.85, 3.52, 4.3, .45, "Affine prior", 22, ORANGE, True, PP_ALIGN.CENTER)
    text(s, 7.85, 4.12, 4.3, .8, "Ridge: causal X → RUL\n경계를 강제하지 않음", 15, INK, True, PP_ALIGN.CENTER)
    pill(s, 9.05, 5.02, 1.9, "PRIOR ON", ORANGE)
    box(s, 2.05, 6.05, 9.2, .65, PALE_GRAY, MUTED)
    text(s, 2.25, 6.12, 8.8, .48, "v1_declared의 neural_safety·FEMTO abstention은 역사 보관이며 Final 9-setting이 아니다.", 12, MUTED, True, PP_ALIGN.CENTER)

    # 6 Equations
    slide_no += 1
    s = blank(prs)
    title(s, "수식으로 보는 PP-X Final", "prior가 tail 방향을 주고, residual은 허용된 범위에서만 수정한다.", slide_no)
    formula(s, .7, 1.45, 3.75, 1.72, "m = h − h_EOL", "현재 health와 고장 경계의 거리", PALE_BLUE, BLUE)
    formula(s, 4.78, 1.45, 3.75, 1.72, "ŷ_BQ = m · softplus(q)", "known boundary branch", PALE_BLUE, BLUE)
    formula(s, 8.86, 1.45, 3.75, 1.72, "ŷ_A = wᵀX + b", "unknown boundary branch", PALE_ORANGE, ORANGE)
    line(s, 2.55, 3.48, 10.75, 3.48, MUTED, 2)
    line(s, 6.65, 3.48, 6.65, 3.92, MUTED, 2, True)
    box(s, 2.2, 4.05, 8.95, 1.65, PALE_GREEN, GREEN)
    text(s, 2.45, 4.3, 8.45, .55, "ŷ = D { y_prior + B(z) · tanh[ r(z) / B(z) ] }", 22, GREEN, True, PP_ALIGN.CENTER)
    text(s, 2.45, 5.0, 8.45, .42, "D: domain decoder  ·  B(z): residual authority  ·  r(z): nonlinear deviation", 11.5, INK, False, PP_ALIGN.CENTER)
    text(s, 1.1, 6.22, 11.15, .43, "frozen prior는 외삽 방향을 유지하고, B(z)는 NN correction의 최대 크기를 통제한다.", 15, INK, True, PP_ALIGN.CENTER)

    # 7 Executor approval
    slide_no += 1
    s = blank(prs)
    title(s, "Executor gate — validation에서 한 번 승인", "contract가 허용한 후보만 같은 validation unit에서 비교한다.", slide_no)
    text(s, .85, 1.42, 11.6, .35, "prior-only   ·   unbounded   ·   bounded   ·   dual-scale   ·   transport   ·   history", 14, MUTED, True, PP_ALIGN.CENTER)
    criteria = [
        (.8, "≥ 2%", "validation gain", BLUE, PALE_BLUE),
        (4.62, "≥ 60%", "physical-unit wins", CYAN, RGBColor(230, 247, 249)),
        (8.44, "≤ 1.10", "worst-unit RMSE ratio", ORANGE, PALE_ORANGE),
    ]
    for x, value, label, accent, fill in criteria:
        box(s, x, 2.05, 3.45, 1.75, fill, accent)
        text(s, x, 2.28, 3.45, .65, value, 26, accent, True, PP_ALIGN.CENTER)
        text(s, x, 3.05, 3.45, .38, label, 12, INK, True, PP_ALIGN.CENTER)
    text(s, 5.95, 3.98, 1.4, .34, "AND", 13, MUTED, True, PP_ALIGN.CENTER)
    box(s, 1.35, 4.55, 4.5, 1.25, PALE_GREEN, GREEN)
    text(s, 1.6, 4.68, 4.0, .35, "PASS", 18, GREEN, True, PP_ALIGN.CENTER)
    text(s, 1.6, 5.1, 4.0, .45, "최소 validation loss route 동결", 12.5, INK, True, PP_ALIGN.CENTER)
    box(s, 7.48, 4.55, 4.5, 1.25, RGBColor(251, 235, 235), RED)
    text(s, 7.73, 4.68, 4.0, .35, "FAIL", 18, RED, True, PP_ALIGN.CENTER)
    text(s, 7.73, 5.1, 4.0, .45, "사전 지정 direct / persistence fallback", 12.5, INK, True, PP_ALIGN.CENTER)
    text(s, 1.15, 6.35, 11.0, .4, "test label · test batch 통계 · test-time 재선택 없음", 15, RED, True, PP_ALIGN.CENTER)

    # 8 Frozen timeline
    slide_no += 1
    s = blank(prs)
    title(s, "학습에서 시험까지 — route는 다시 열리지 않는다", "test에는 frozen forward만 남는다.", slide_no)
    stages = [
        (1.0, "TRAIN", "prior·residual\n후보 학습", BLUE),
        (4.1, "VALIDATION", "executor 승인\n2% · 60% · 1.10", ORANGE),
        (7.3, "FREEZE", "route 1개\n파라미터 고정", GREEN),
        (10.4, "TEST", "개별 causal X\nforward 1회", INK),
    ]
    line(s, 1.65, 3.65, 11.05, 3.65, RGBColor(195, 203, 211), 5)
    for x, name, body, accent in stages:
        dot(s, x + .42, 3.32, .66, accent)
        pill(s, x, 2.05, 1.5, name, accent, WHITE, 10)
        text(s, x - .35, 4.22, 2.2, .85, body, 13, INK, True, PP_ALIGN.CENTER)
    box(s, 3.1, 5.75, 7.15, .68, PALE_ORANGE, ORANGE)
    text(s, 3.3, 5.86, 6.75, .45, "금지: test outcome으로 route 수정 · calibration · 재학습", 13, RED, True, PP_ALIGN.CENTER)

    # 9 Portfolio map
    slide_no += 1
    s = blank(prs)
    title(s, "9-setting Final portfolio", "하나의 거대 network가 아니라, 같은 실행 계약을 따르는 typed routes다.", slide_no)
    box(s, 4.75, 1.33, 3.85, .72, INK, INK)
    text(s, 4.75, 1.33, 3.85, .72, "PP-X FINAL CONTRACT", 16, WHITE, True, PP_ALIGN.CENTER)
    groups = [
        (.65, 2.85, 3.6, 2.7, "BQ branch", "Sunwoda · RWTH\nMICH", BLUE, PALE_BLUE,
         "bounded · dual-scale"),
        (4.85, 2.85, 3.6, 2.7, "Affine branch", "Virkler · MATR19", ORANGE, PALE_ORANGE,
         "gated residual · calibration"),
        (9.05, 2.85, 3.6, 2.7, "Affine + adapter", "HUST · MATR-b2\nNASA · N-CMAPSS", GREEN, PALE_GREEN,
         "transport · history"),
    ]
    for x, y, w, h, name, datasets, accent, fill, executor in groups:
        line(s, 6.68, 2.05, x + w/2, y, accent, 2.5, True)
        box(s, x, y, w, h, fill, accent)
        text(s, x + .2, y + .22, w - .4, .42, name, 18, accent, True, PP_ALIGN.CENTER)
        text(s, x + .2, y + .86, w - .4, .68, datasets, 14, INK, True, PP_ALIGN.CENTER)
        pill(s, x + .42, y + 1.93, w - .84, executor, accent, WHITE, 10.5)
    text(s, 1.0, 6.2, 11.3, .42, "공통점: prior–residual 역할 분리 + validation-approved executor + frozen test execution", 14.5, INK, True, PP_ALIGN.CENTER)

    # 10 scores
    slide_no += 1
    s = blank(prs)
    title(s, "성능 — 9개 setting 모두 R² > 0", "retrospective main portfolio · pooled R²", slide_no)
    add_picture(s, charts["scores"], .72, 1.28, 11.35, 5.45)
    pill(s, 10.7, 6.42, 1.65, "macro 0.807", BLUE)

    # 11 equal budget
    slide_no += 1
    s = blank(prs)
    title(s, "동일 탐색 예산 비교", "각 setting에서 30-candidate strongest baseline과 비교", slide_no)
    add_picture(s, charts["equal"], .72, 1.25, 11.5, 5.35)
    box(s, 9.0, 5.95, 3.35, .72, PALE_BLUE, BLUE)
    text(s, 9.2, 6.04, 2.95, .5, "8 / 9 우세  ·  p = .039", 15, BLUE, True, PP_ALIGN.CENTER)
    text(s, .85, 6.08, 7.7, .45, "예외: Virkler 0.888 vs FT-Transformer 0.890", 12, MUTED, True)

    # 12 core ablation
    slide_no += 1
    s = blank(prs)
    title(s, "Core ablation — prior만으로는 부족하다", "동일 row · seeds 42–46 prediction ensemble", slide_no)
    add_picture(s, charts["core"], .72, 1.33, 8.55, 5.25)
    box(s, 9.35, 1.7, 3.15, 3.95, PALE_BLUE, BLUE)
    text(s, 9.65, 2.0, 2.55, .48, "Nonlinear residual", 17, BLUE, True, PP_ALIGN.CENTER)
    text(s, 9.65, 2.78, 2.55, 1.5, "Sunwoda  +0.659\nRWTH       +0.220\nMICH        +3.811", 15, INK, True)
    line(s, 9.7, 4.6, 12.15, 4.6, BLUE, 2)
    text(s, 9.62, 4.83, 2.6, .55, "세 setting 모두\nunit-level 강한 근거", 11.5, MUTED, True, PP_ALIGN.CENTER)

    # 13 executor ablation
    slide_no += 1
    s = blank(prs)
    title(s, "Executor ablation — 효과 부호가 바뀐다", "optional module을 전역 default로 켜면 안 되는 직접 근거", slide_no)
    add_picture(s, charts["executor"], .72, 1.3, 9.2, 5.4)
    box(s, 9.95, 1.65, 2.65, 1.55, PALE_GREEN, GREEN)
    text(s, 10.15, 1.84, 2.25, .36, "켜야 하는 경우", 14, GREEN, True, PP_ALIGN.CENTER)
    text(s, 10.15, 2.34, 2.25, .48, "HUST transport\nMICH dual-scale", 11.5, INK, True, PP_ALIGN.CENTER)
    box(s, 9.95, 3.55, 2.65, 1.55, RGBColor(251, 235, 235), RED)
    text(s, 10.15, 3.74, 2.25, .36, "끄는 반례", 14, RED, True, PP_ALIGN.CENTER)
    text(s, 10.15, 4.24, 2.25, .48, "RWTH dual-scale\nMICH fixed bound", 11.5, INK, True, PP_ALIGN.CENTER)
    text(s, 9.92, 5.67, 2.7, .58, "그래서 contract + Val\n승인이 필요하다.", 12.5, INK, True, PP_ALIGN.CENTER)

    # 14 statistics
    slide_no += 1
    s = blank(prs)
    title(s, "통계검정 — 행이 아니라 물리 unit이 표본", "요약 효과와 개별 component evidence를 분리한다.", slide_no)
    metrics = [
        (.75, "9 / 9", "vs matched direct\n방향 일치", BLUE, PALE_BLUE),
        (3.9, "p = .0039", "domain exact\nsign test", CYAN, RGBColor(230, 247, 249)),
        (7.05, "33.8%", "geometric-mean\nRMSE 감소", GREEN, PALE_GREEN),
        (10.2, "[0.17, 0.67]", "hierarchical CI\nlog-RMSE gain", ORANGE, PALE_ORANGE),
    ]
    for x, value, label, accent, fill in metrics:
        box(s, x, 1.55, 2.4, 2.0, fill, accent)
        text(s, x, 1.82, 2.4, .6, value, 22, accent, True, PP_ALIGN.CENTER)
        text(s, x, 2.58, 2.4, .65, label, 11.5, INK, True, PP_ALIGN.CENTER)
    line(s, 1.0, 4.1, 12.3, 4.1, RGBColor(220, 225, 230), 1.2)
    text(s, .9, 4.45, 2.5, .4, "개별 BH 근거", 15, INK, True)
    pill(s, 3.0, 4.47, 1.55, "Sunwoda", BLUE)
    pill(s, 4.75, 4.47, 1.55, "MATR-b2", GREEN)
    text(s, .9, 5.25, 2.5, .4, "해석 제한", 15, INK, True)
    text(s, 3.0, 5.13, 8.8, .8, "5 seeds는 독립 cohort가 아님 · NASA/N-CMAPSS unit 수가 작음\nretrospective split은 prospective 확증을 대체하지 않음", 12.5, MUTED, True)
    box(s, .9, 6.25, 11.7, .55, RGBColor(251, 235, 235), RED)
    text(s, 1.15, 6.31, 11.2, .42, "FEMTO v1_declared abstention은 Final BH family에서 제외", 12, RED, True, PP_ALIGN.CENTER)

    # 15 claim boundary
    slide_no += 1
    s = blank(prs)
    title(s, "무엇을 주장하고, 무엇을 주장하지 않는가", "성과보다 실행 계약의 범위를 먼저 고정한다.", slide_no)
    box(s, .72, 1.45, 5.75, 4.85, PALE_GREEN, GREEN)
    text(s, 1.05, 1.75, 5.1, .45, "말할 수 있음", 21, GREEN, True)
    text(s, 1.05, 2.52, 5.05, 2.8,
         "✓ 9-setting summary gain\n\n✓ prior–residual 역할 분리\n\n✓ executor 효과의 반례\n\n✓ test 전 route 동결", 16, INK, True)
    box(s, 6.86, 1.45, 5.75, 4.85, RGBColor(251, 235, 235), RED)
    text(s, 7.19, 1.75, 5.1, .45, "말하면 안 됨", 21, RED, True)
    text(s, 7.19, 2.52, 5.05, 2.8,
         "× 모든 데이터에서 SOTA\n\n× OOF mode-stability가 Final에서 실행됨\n\n× 모든 executor가 항상 유효\n\n× 미래 cohort 성능 보장", 16, INK, True)

    # 16 takeaway
    slide_no += 1
    s = blank(prs, True)
    pill(s, .72, .62, 1.7, "TAKEAWAY", BLUE)
    text(s, .82, 1.45, 11.6, 1.25, "경계가 있으면 BQ,\n없으면 affine.", 31, WHITE, True)
    text(s, .84, 3.13, 11.2, .55, "그 위의 residual과 executor만 validation에서 승인한다.", 21, RGBColor(214, 224, 234), True)
    line(s, .84, 4.15, 12.2, 4.15, BLUE, 4)
    steps = [("1", "Prior", "BQ / affine"), ("2", "Fit", "bounded residual"), ("3", "Approve", "2% · 60% · 1.10"), ("4", "Freeze", "test forward")]
    for i, (num, name, body) in enumerate(steps):
        x = .85 + i * 3.02
        dot(s, x, 4.72, .52, BLUE if i < 3 else GREEN)
        text(s, x, 4.72, .52, .52, num, 12, WHITE, True, PP_ALIGN.CENTER)
        text(s, x + .68, 4.62, 2.15, .38, name, 14, WHITE, True)
        text(s, x + .68, 5.03, 2.15, .35, body, 10.5, RGBColor(180, 195, 210))
    text(s, .82, 6.35, 4.5, .45, "박진서 · Q&A", 15, WHITE, True)
    text(s, 11.4, 6.75, 1.2, .3, "16 / 16", 10, RGBColor(180, 195, 210), True, PP_ALIGN.RIGHT)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Saved {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
