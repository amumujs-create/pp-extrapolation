#!/usr/bin/env python3
"""Align the portfolio table labels with the English chart labels."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt


ROOT = Path(__file__).resolve().parents[2]
DECKS = (
    ROOT / "ppt" / "PP-X_Final_Research_Detailed_v2.pptx",
    ROOT / "ppt" / "PP-X_Final_Research_20min.pptx",
)
LABELS = {
    "화중 배터리": "HUST\n화중 배터리",
    "선우다 상용셀": "Sunwoda\n선우다 상용셀",
    "항공기 엔진": "N-CMAPSS\n항공기 엔진",
    "알루미늄 균열": "Virkler\n알루미늄 균열",
    "아헨 배터리": "RWTH\n아헨 배터리",
    "MIT 배치2": "MATR batch2\nMIT 배치2",
    "미시간 배터리": "MICH\n미시간 배터리",
    "NASA 실험셀": "NASA PCoE\nNASA 실험셀",
    "MIT 2019": "MATR 2019\nMIT 2019",
}


def style_cell(cell) -> None:
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for paragraph in cell.text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.LEFT
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
            rpr = run._r.get_or_add_rPr()
            east_asian = rpr.find(qn("a:ea"))
            if east_asian is not None:
                east_asian.set("typeface", "AppleGothic")


def update(path: Path) -> int:
    prs = Presentation(path)
    changed = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            table = shape.table
            if len(table.columns) != 3:
                continue
            headers = [table.cell(0, column).text.strip() for column in range(3)]
            if headers != ["셋", "Executor", "R²"]:
                continue
            for row in range(1, len(table.rows)):
                cell = table.cell(row, 0)
                current = cell.text.strip()
                if current in LABELS:
                    cell.text = LABELS[current]
                    style_cell(cell)
                    changed += 1
    prs.save(path)
    return changed


def main() -> None:
    for path in DECKS:
        print(f"{path.name}: {update(path)} labels aligned")


if __name__ == "__main__":
    main()
