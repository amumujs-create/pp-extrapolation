#!/usr/bin/env python3
"""Append the complete PP-X Final inference appendix to the detailed deck."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt


ROOT = Path(__file__).resolve().parents[2]
DECK = ROOT / "ppt" / "PP-X_Final_Research_Detailed_v2.pptx"
PX = 9525
C = {
    "ink": RGBColor(0x22, 0x22, 0x22),
    "muted": RGBColor(0x66, 0x66, 0x66),
    "rule": RGBColor(0xE8, 0xE8, 0xE8),
    "blue": RGBColor(0x00, 0x72, 0xB2),
    "orange": RGBColor(0xE6, 0x9F, 0x00),
    "soft": RGBColor(0xF5, 0xF5, 0xF5),
    "soft_blue": RGBColor(0xE8, 0xF1, 0xF8),
    "soft_orange": RGBColor(0xFD, 0xF4, 0xE3),
    "white": RGBColor(0xFF, 0xFF, 0xFF),
    "red": RGBColor(0x9B, 0x1C, 0x1C),
}


def px(value: float) -> int:
    return int(value * PX)


def set_run(run, size=18, color=None, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Arial"
    run.font.color.rgb = color or C["ink"]
    rpr = run._r.get_or_add_rPr()
    east_asian = rpr.find(qn("a:ea"))
    if east_asian is None:
        from lxml import etree
        east_asian = etree.SubElement(rpr, qn("a:ea"))
    east_asian.set("typeface", "AppleGothic")


def add_text(slide, left, top, width, height, value, size=18, color=None,
             bold=False, align="left"):
    shape = slide.shapes.add_textbox(px(left), px(top), px(width), px(height))
    tf = shape.text_frame
    tf.word_wrap = True
    for index, value_line in enumerate(str(value).split("\n")):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.alignment = {
            "left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT
        }[align]
        run = paragraph.add_run()
        run.text = value_line
        set_run(run, size, color or C["ink"], bold)
    return shape


def rect(slide, left, top, width, height, fill=None, border=None, rounded=False):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if rounded else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, px(left), px(top), px(width), px(height))
    shape.line.fill.background()
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if border is not None:
        shape.line.color.rgb = border
        shape.line.width = Pt(1.25)
    return shape


def hline(slide, x1, x2, y, color):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x1), px(y), px(x2), px(y))
    shape.line.color.rgb = color
    shape.line.width = Pt(1.25)
    return shape


def add_table(slide, left, top, width, height, headers, rows, *, font_size=11):
    table = slide.shapes.add_table(
        1 + len(rows), len(headers), px(left), px(top), px(width), px(height)
    ).table
    for column in table.columns:
        column.width = px(width / len(headers))

    def fill_cell(row, column, value, fill, color, bold, align):
        cell = table.cell(row, column)
        cell.text = ""
        paragraph = cell.text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.LEFT if align == "left" else PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = str(value)
        set_run(run, font_size, color, bold)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill

    for column, value in enumerate(headers):
        fill_cell(0, column, value, C["ink"], C["white"], True, "center")
    for row, values in enumerate(rows, 1):
        fill = C["soft"] if row % 2 == 0 else C["white"]
        for column, value in enumerate(values):
            fill_cell(row, column, value, fill, C["ink"], column == 0, "left" if column == 0 else "center")
    return table


def head(slide, heading, sub=""):
    rect(slide, 0, 0, 1280, 2, C["ink"])
    add_text(slide, 48, 22, 1180, 32, heading, 22, C["ink"], True)
    if sub:
        add_text(slide, 48, 54, 1180, 20, sub, 12, C["muted"])


def blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = C["white"]
    background.line.fill.background()
    return slide


def footer(slide, number: int, total: int) -> None:
    add_text(slide, 1140, 674, 88, 18, f"{number}/{total}", 9, C["muted"], True, "right")


def build() -> None:
    prs = Presentation(DECK)
    original = len(prs.slides)
    total = original + 5

    # 1. Statistics map
    s = blank(prs)
    head(s, "통계검정 지도 — 질문마다 검정이 다르다",
         "내부 구조·외부 경쟁력·component·붕괴 안정성을 하나의 p값으로 섞지 않는다.")
    add_table(
        s, 48, 95, 1184, 390,
        ["질문", "비교 대상", "표본 단위", "검정", "판정"],
        [
            ["Prior 구조가 필요한가?", "PP-X vs matched direct", "9 setting · 77 physical units",
             "exact sign + hierarchical bootstrap", "9/9 · p=.0039 · RMSE −33.8%"],
            ["강한 ML보다 경쟁력 있나?", "PP-X vs 동일예산 최강 baseline", "9 settings",
             "two-sided exact sign", "8/9 · p=.0391"],
            ["각 module이 필요한가?", "component on vs off", "physical units",
             "unit sign/boot + 24개 BH", "일부 개선·일부 악화 유의"],
            ["외삽이 덜 무너지나?", "9모델의 setting별 R² 분포", "9 settings",
             "MAD permutation + SD bootstrap", "관찰상 우수; Holm 후 전역 비유의"],
            ["Seed가 더 안정적인가?", "5-seed refits", "optimizer seeds 5개",
             "exact sign/기술통계", "p<.05 확증 불가"],
        ],
        font_size=10,
    )
    rect(s, 120, 525, 1040, 72, C["soft_blue"], C["blue"], False)
    add_text(s, 140, 538, 1000, 46,
             "핵심 주장: 현재 9개 retrospective 외삽 setting에서 PP-X는 prior 없는 direct보다 "
             "평균 오차가 작고, 강한 동일예산 비교군보다 8/9에서 높으며, 심한 음의 R² 붕괴가 덜 관찰됐다.",
             12, C["ink"], True, "center")
    footer(s, original + 1, total)

    # 2. Matched direct and hierarchical bootstrap
    s = blank(prs)
    head(s, "검정 1 — prior–residual 구조의 평균 효과",
         "FT-Transformer 비교가 아니다. 같은 데이터·split에서 prior 구조를 제거한 내부 대조군과 비교한다.")
    rect(s, 48, 100, 350, 410, C["soft_blue"], C["blue"], False)
    add_text(s, 68, 118, 310, 28, "비교 목적", 15, C["blue"], True, "center")
    add_text(s, 72, 170, 302, 150,
             "PP-X Final\nvs\nmatched direct\n\n"
             "같은 입력·split의 direct 예측 모델에서\nBQ/affine prior-residual 구조만 제거",
             13, C["ink"], True, "center")
    add_text(s, 72, 370, 302, 92,
             "묻는 질문\n“일반 direct 예측보다\nprior 구조가 실제로 도움 되는가?”",
             12, C["muted"], True, "center")

    rect(s, 430, 100, 380, 410, C["soft"], C["ink"], False)
    add_text(s, 450, 118, 340, 28, "계층 bootstrap", 15, C["ink"], True, "center")
    add_text(s, 452, 170, 336, 250,
             "① 각 물리 unit의 RMSE 비교\n\n"
             "② 데이터셋을 동일 확률로 재표집\n\n"
             "③ 선택된 데이터셋 안에서 unit 재표집\n\n"
             "④ 반복마다 log-RMSE 이득 계산\n\n"
             "→ 큰 데이터셋이 결론을 독점하지 않음",
             12, C["ink"])
    add_text(s, 452, 442, 336, 42, "9 settings · 77 physical units", 12, C["muted"], True, "center")

    rect(s, 842, 100, 390, 410, C["soft_blue"], C["blue"], False)
    add_text(s, 862, 118, 350, 28, "결과와 해석", 15, C["blue"], True, "center")
    add_text(s, 862, 172, 350, 56, "9 / 9 우세  ·  p=.0039", 18, C["blue"], True, "center")
    hline(s, 875, 1198, 250, C["blue"])
    add_text(s, 862, 270, 350, 54, "평균 RMSE 33.8% 감소", 17, C["blue"], True, "center")
    add_text(s, 862, 328, 350, 42, "95% CI: 15.8%–48.6%", 13, C["ink"], True, "center")
    add_text(s, 862, 394, 350, 76,
             "log CI [.172,.665]를\n사람이 읽기 쉬운 RMSE 감소율로 환산",
             11, C["muted"], False, "center")
    add_text(s, 55, 540, 1170, 50,
             "말할 수 있음: 현재 평가 범위에서 prior–residual 구조의 평균 오차 감소가 0보다 크다. "
             "말할 수 없음: 모든 MLP·Transformer보다 항상 우수하다.",
             12, C["ink"], True, "center")
    footer(s, original + 2, total)

    # 3. Equal-budget external comparison
    s = blank(prs)
    head(s, "검정 2 — 동일예산 강력 baseline과의 외부 경쟁력",
         "각 setting에서 30개 후보를 탐색한 가장 강한 baseline과 PP-X Final을 비교한다.")
    add_table(
        s, 48, 92, 1184, 380,
        ["Setting", "PP-X", "동일예산 최강 baseline", "차이", "승패"],
        [
            ["HUST", ".958", "GroupDRO .955", "+.003", "승"],
            ["Sunwoda", ".939", "RBF .838", "+.102", "승"],
            ["N-CMAPSS", ".937", "Engression .932", "+.005", "승"],
            ["Virkler", ".888", "FT-Transformer .890", "−.002", "패"],
            ["RWTH", ".878", "RBF .732", "+.146", "승"],
            ["MATR-b2", ".862", "MLP .813", "+.049", "승"],
            ["MICH", ".751", "monotone −.686", "+1.437", "승"],
            ["NASA", ".584", "Engression .583", "~0", "승"],
            ["MATR19", ".466", "FT-Transformer .342", "+.123", "승"],
        ],
        font_size=10,
    )
    rect(s, 95, 505, 500, 85, C["soft_blue"], C["blue"], False)
    add_text(s, 110, 518, 470, 54, "8 / 9 우세 · two-sided exact p=.0391", 15, C["blue"], True, "center")
    rect(s, 680, 505, 500, 85, C["soft_orange"], C["orange"], False)
    add_text(s, 695, 518, 470, 54, "Virkler에서는 FT-Transformer가 +.002 우세", 13, C["orange"], True, "center")
    add_text(s, 60, 615, 1160, 34,
             "이 검정은 prior 구조의 내부 효과가 아니라, 동일한 탐색 예산 아래 외부 모델과의 경쟁력을 묻는다.",
             12, C["muted"], True, "center")
    footer(s, original + 3, total)

    # 4. Component inference and final-route matching
    s = blank(prs)
    head(s, "검정 3 — component on/off와 Final route의 일치",
         "강제 on/off ablation을 물리 unit 단위로 비교하고 24개 동시검정에 BH 보정을 적용했다.")
    add_table(
        s, 48, 95, 1184, 340,
        ["Component", "Setting", "OFF → ON R²", "BH q", "Ablation 판정", "실제 Final"],
        [
            ["nonlinear residual", "Sunwoda", ".281 → .939", ".0188", "개선", "ON"],
            ["nonlinear residual", "RWTH", ".659 → .878", ".0208", "개선", "ON"],
            ["nonlinear residual", "MICH", "−3.343 → .468", ".0208", "개선", "ON"],
            ["dual-scale", "MICH", ".468 → .751", ".0313", "개선", "ON"],
            ["dual-scale", "RWTH", ".878 → .842", ".0208", "유의하게 악화", "OFF · fixed .878"],
            ["fixed bound", "MICH", ".759 → .468", ".0313", "유의하게 악화", "OFF · dual .751"],
            ["regime transport", "HUST", ".829 → .958", ".0029", "개선", "ON"],
            ["regime transport", "MATR-b2", ".675 → .862", ".0188", "개선", "ON"],
        ],
        font_size=9.5,
    )
    rect(s, 70, 480, 1140, 86, C["soft_blue"], C["blue"], False)
    add_text(s, 90, 493, 1100, 58,
             "실험과 실제 결과가 매칭됨: MICH는 dual-scale을 켰고, RWTH는 dual-scale을 끄고 fixed bound를 유지했다. "
             "음성 ablation은 모델 실패가 아니라 잘못된 전역 적용을 막는 근거다.",
             12, C["ink"], True, "center")
    add_text(s, 70, 605, 1140, 44,
             "BH 보정 목적: 24개 비교 중 우연히 유의한 결과가 생기는 허위발견률을 통제.",
             11.5, C["muted"], True, "center")
    footer(s, original + 4, total)

    # 5. Collapse resistance and statistical boundary
    s = blank(prs)
    head(s, "검정 4 — ‘외삽이 덜 무너진다’의 근거와 한계",
         "성능 평균뿐 아니라 setting 간 편차·최악 점수·음의 R² 발생을 함께 본다.")
    add_table(
        s, 48, 96, 730, 330,
        ["모델", "macro R²", "setting SD", "최악 R²", "R²>0"],
        [
            ["PP-X", ".807", ".174", ".466", "9/9"],
            ["Engression", ".257", ".949", "−1.580", "7/9"],
            ["FT-Transformer", ".115", "1.005", "−2.010", "7/9"],
            ["plain MLP", ".093", "1.007", "−2.140", "6/9"],
            ["GroupDRO", ".011", ".998", "−2.057", "5/9"],
        ],
        font_size=11,
    )
    rect(s, 820, 96, 410, 150, C["soft_blue"], C["blue"], False)
    add_text(s, 840, 112, 370, 28, "관찰된 붕괴 완화", 15, C["blue"], True, "center")
    add_text(s, 840, 154, 370, 70, "가장 작은 SD\n가장 높은 최악 R²\n유일한 9/9 양의 R²", 13, C["ink"], True, "center")
    rect(s, 820, 276, 410, 150, C["soft_orange"], C["orange"], False)
    add_text(s, 840, 292, 370, 28, "통계적 경계", 15, C["orange"], True, "center")
    add_text(s, 840, 334, 370, 70, "Engression 대비 MAD p=.0391\n8모델 Holm 보정 q=.2734\n→ 전 모델 대비 확증은 아님", 12, C["ink"], True, "center")
    rect(s, 60, 475, 1160, 92, C["soft_blue"], C["blue"], False)
    add_text(s, 85, 489, 1110, 62,
             "허용 주장: “현재 9개 retrospective 외삽 setting에서 PP-X는 심한 성능 붕괴가 덜 관찰됐다.”\n"
             "금지 주장: “모든 미래 데이터와 모든 모델에 대해 안정성 우월이 통계적으로 확증됐다.”",
             12.5, C["ink"], True, "center")
    add_text(s, 70, 610, 1140, 35,
             "5-seed는 optimizer 재학습 기술통계다. n=5 양측 exact 검정 최소 p=.0625이므로 독립적인 안정성 확증으로 쓰지 않는다.",
             11.5, C["muted"], True, "center")
    footer(s, original + 5, total)

    prs.save(DECK)
    print(f"Updated {DECK} ({original} -> {len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
