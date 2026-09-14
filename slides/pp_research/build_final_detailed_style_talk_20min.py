#!/usr/bin/env python3
"""Create a 20-minute talk by selecting slides from the Final detailed deck."""
from __future__ import annotations

import re
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

from append_final_detailed_statistics import C, add_table, add_text, blank, head, rect


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "ppt" / "PP-X_Final_Research_Detailed_v2.pptx"
OUTPUT = ROOT / "ppt" / "PP-X_Final_Research_20min.pptx"
COMPETITOR_FIGURE = ROOT.parent / "ppt" / "pp" / "_build" / "figs" / "competitor_bars.png"

# Narrative: why extrapolation → PP-X/PAE route → Final algorithm → results
# → fair comparison → complete inference explanation → Q&A.
KEEP = (1, 2, 3, 4, 6, 7, 8, 9, 10, 11, 14, 43, 44, 46, 47, 40)


def remove_slide(prs: Presentation, zero_based_index: int) -> None:
    slide_id = prs.slides._sldIdLst[zero_based_index]
    relationship_id = slide_id.rId
    prs.part.drop_rel(relationship_id)
    del prs.slides._sldIdLst[zero_based_index]


def replace_text(slide, old: str, new: str) -> None:
    for shape in slide.shapes:
        if not hasattr(shape, "text_frame"):
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)


def build() -> None:
    shutil.copy2(SOURCE, OUTPUT)
    prs = Presentation(OUTPUT)
    selected = set(KEEP)
    for index in range(len(prs.slides), 0, -1):
        if index not in selected:
            remove_slide(prs, index - 1)

    retained_ids = list(prs.slides._sldIdLst)
    by_source_index = dict(zip(sorted(KEEP), retained_ids))

    formula_slide = blank(prs)
    head(formula_slide, "수식 1 — Final prior는 BQ 또는 affine",
         "known_boundary가 prior 함수형만 고른다. 둘 다 ON이며 residual 후보 학습으로 이어진다.")
    rect(formula_slide, 55, 105, 550, 400, C["soft_blue"], C["blue"], True)
    add_text(formula_slide, 80, 130, 500, 34, "known_boundary = TRUE  →  BQ", 17, C["blue"], True, "center")
    add_text(formula_slide, 80, 205, 500, 58, "m = health − EOL boundary", 18, C["ink"], True, "center")
    add_text(formula_slide, 80, 278, 500, 62, "ŷprior = m · softplus(q)", 21, C["blue"], True, "center")
    add_text(formula_slide, 90, 370, 480, 78,
             "경계거리 m을 곱하므로 EOL 경계에서 RUL=0.\nq는 단위 경계거리당 남은 수명.",
             12, C["muted"], False, "center")
    rect(formula_slide, 675, 105, 550, 400, C["soft_orange"], C["orange"], True)
    add_text(formula_slide, 700, 130, 500, 34, "known_boundary = FALSE  →  affine", 17, C["orange"], True, "center")
    add_text(formula_slide, 700, 225, 500, 62, "ŷprior = wᵀX + b", 22, C["orange"], True, "center")
    add_text(formula_slide, 710, 340, 480, 90,
             "causal feature X에서 RUL 기본 추세를 Ridge로 학습.\n경계 0을 강제하지 않는 비경계형 tail.",
             12, C["muted"], False, "center")
    rect(formula_slide, 175, 545, 930, 70, C["soft"], C["ink"], True)
    add_text(formula_slide, 195, 558, 890, 42,
             "Final prior gate: BQ냐 affine이냐만 결정 · OOF/group/mode 검사 없음 · prior_weight=1",
             13, C["ink"], True, "center")
    formula_id = prs.slides._sldIdLst[-1]

    residual_slide = blank(prs)
    head(residual_slide, "수식 2 — residual 후보와 executor의 역할",
         "C2는 bounded 하나만 학습하지 않는다. 후보군을 학습하고 C3 validation이 하나를 고른다.")
    add_table(
        residual_slide, 45, 96, 1190, 400,
        ["후보", "예측식 / 동작", "무엇을 바꾸나", "Final에서의 위치"],
        [
            ["prior-only", "D{yprior}", "residual 없음", "가장 단순한 기준"],
            ["unbounded", "D{yprior + r(z)}", "수정폭 제한 없음", "공통 residual 후보"],
            ["fixed bounded", "D{yprior + B·tanh[r(z)/B]}", "|수정|≤B", "Sunwoda·RWTH"],
            ["dual-scale", "D{yprior + B(z)·tanh[r(z)/B(z)]}", "support별 local/broad bound", "MICH만 ON"],
            ["history", "z = adapter(current, causal history)", "residual 입력 시간척도", "NASA·N-CMAPSS"],
            ["transport", "ŷ′ = ag·ŷ + cg", "regime 출력 scale·offset", "HUST·MATR-b2"],
            ["direct fallback", "D{fNN(X)}", "prior-residual 미사용", "Executor Val FAIL"],
        ],
        font_size=9.5,
    )
    rect(residual_slide, 100, 535, 1080, 82, C["soft_blue"], C["blue"], True)
    add_text(residual_slide, 125, 548, 1030, 54,
             "C2 Fit = prior-only·unbounded·bounded 후보 학습 → C3 Approve = 2%·60%·1.10 통과 executor 1개 동결",
             12.5, C["ink"], True, "center")
    residual_id = prs.slides._sldIdLst[-1]

    comparison_slide = blank(prs)
    head(comparison_slide, "동일예산 비교 — 그림과 검정의 의미",
         "왼쪽 heatmap은 9 settings × 9 methods의 pooled R². 오른쪽은 이 그림에서 계산한 검정이다.")
    comparison_slide.shapes.add_picture(
        str(COMPETITOR_FIGURE), 20 * 9525, 85 * 9525, 820 * 9525, 520 * 9525
    )
    rect(comparison_slide, 865, 100, 365, 145, C["soft_blue"], C["blue"], True)
    add_text(comparison_slide, 885, 114, 325, 26, "무엇을 비교했나", 14, C["blue"], True, "center")
    add_text(comparison_slide, 885, 152, 325, 74,
             "각 setting에서 후보 30개를 탐색한\n가장 강한 baseline vs PP-X Final",
             11.5, C["ink"], True, "center")
    rect(comparison_slide, 865, 275, 365, 125, C["soft_blue"], C["blue"], True)
    add_text(comparison_slide, 885, 288, 325, 35, "8 / 9 우세 · p=.0391", 16, C["blue"], True, "center")
    add_text(comparison_slide, 885, 334, 325, 48,
             "9개 setting 승패의\ntwo-sided exact sign test", 10.5, C["muted"], False, "center")
    rect(comparison_slide, 865, 430, 365, 125, C["soft_orange"], C["orange"], True)
    add_text(comparison_slide, 885, 443, 325, 30, "예외 · Virkler", 14, C["orange"], True, "center")
    add_text(comparison_slide, 885, 482, 325, 54,
             "PP-X .888\nFT-Transformer .890 (+.002)", 11.5, C["ink"], True, "center")
    add_text(comparison_slide, 865, 585, 365, 45,
             "matched direct 9/9 검정과는 별개", 11, C["red"], True, "center")
    comparison_id = prs.slides._sldIdLst[-1]

    for slide_id in retained_ids:
        prs.slides._sldIdLst.remove(slide_id)
    for slide_id in (formula_id, residual_id, comparison_id):
        prs.slides._sldIdLst.remove(slide_id)
    narrative = (
        1, 2, 3, 4, 6, 7, 8, "formula", "residual", 9, 10, 11, 14,
        "comparison", 43, 44, 46, 47, 40,
    )
    custom = {"formula": formula_id, "residual": residual_id, "comparison": comparison_id}
    for item in narrative:
        prs.slides._sldIdLst.append(custom[item] if isinstance(item, str) else by_source_index[item])

    total = len(prs.slides)
    if total != 19:
        raise RuntimeError(f"expected 19 slides, got {total}")

    # Make the copied cover and method wording specific to this talk.
    replace_text(prs.slides[0], "전체 연구 지도", "20분 요약 발표")
    for slide in prs.slides:
        replace_text(
            slide,
            "route가 승인되지 않으면 해당 contract에 미리 적은 direct·persistence fallback 또는 abstention으로 간다.",
            "Final에서 executor가 승인되지 않으면 contract에 미리 적은 direct·persistence fallback으로 간다.",
        )

    for page, slide in enumerate(prs.slides, 1):
        found = False
        for shape in slide.shapes:
            if not hasattr(shape, "text_frame"):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if re.fullmatch(r"\d+/\d+", run.text.strip()):
                        run.text = f"{page}/{total}"
                        found = True
        # Appended statistics slides can store footer text without a separate
        # run. Add the same restrained footer number when none was detected.
        if page not in (1, total) and not found:
            shape = slide.shapes.add_textbox(1140 * 9525, 674 * 9525, 88 * 9525, 18 * 9525)
            paragraph = shape.text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.RIGHT
            run = paragraph.add_run()
            run.text = f"{page}/{total}"
            run.font.name = "Arial"
            run.font.size = Pt(9)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    prs.save(OUTPUT)
    print(f"Saved {OUTPUT} ({total} slides)")


if __name__ == "__main__":
    build()
