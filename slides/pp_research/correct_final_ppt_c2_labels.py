#!/usr/bin/env python3
"""Correct C2 labels: C2 fits a residual candidate family, not only bounded."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation


ROOT = Path(__file__).resolve().parents[2]
DECKS = (
    ROOT / "ppt" / "PP-X_Final_Research_Detailed_v2.pptx",
    ROOT / "ppt" / "PP-X_Final_Research_20min.pptx",
)
REPLACEMENTS = {
    "prior + bounded\nresidual 학습": "prior + residual\n후보 학습",
    "weak prior + bounded residual\nValidation evidence로 executor 승인 · 실패 시 fallback":
        "weak prior + residual candidate family\nValidation evidence로 executor 승인 · 실패 시 fallback",
    "Prior ON일 때만\nŷ ≈ prior + bounded residual\nNN이 prior를 뒤집지 못하게 수정폭 제한":
        "Prior ON일 때\nprior-only · unbounded · bounded 후보 학습\nbound 적용 여부는 C3가 선택",
    "prior 기본 경로 고정 + bounded residual 학습 (과도한 수정 차단)":
        "prior 기본 경로 + residual 후보군 학습 (bound는 선택형)",
}


def replace_in_shape(shape, old: str, new: str) -> bool:
    if not hasattr(shape, "text_frame") or old not in shape.text:
        return False
    # Replacing the whole text frame preserves the shape style and avoids
    # partial-run misses when a phrase spans multiple PowerPoint runs.
    if shape.text == old:
        paragraphs = shape.text_frame.paragraphs
        lines = new.split("\n")
        if len(paragraphs) == len(lines) and all(p.runs for p in paragraphs):
            for paragraph, line in zip(paragraphs, lines):
                paragraph.runs[0].text = line
                for run in paragraph.runs[1:]:
                    run.text = ""
        else:
            shape.text = new
        return True
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)
                return True
    # Multi-run fallback. Formatting within this explanatory box is uniform.
    shape.text = shape.text.replace(old, new)
    return True


def correct(path: Path) -> int:
    prs = Presentation(path)
    changed = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            for old, new in REPLACEMENTS.items():
                changed += int(replace_in_shape(shape, old, new))
    prs.save(path)
    return changed


def main() -> None:
    for path in DECKS:
        print(f"{path.name}: {correct(path)} labels corrected")


if __name__ == "__main__":
    main()
